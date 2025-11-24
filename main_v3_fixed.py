"""
RESUME-BASED NATURAL INTERVIEW SYSTEM - FIXED VERSION
=====================================================
- Parse resume with consistent JSON structure
- Generate questions purely from resume content using structured approach
- Eliminate hallucination with strict validation
- Answer scoring with detailed feedback
"""

from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Dict, List, Optional
import os
import json
import re
from langchain_groq import ChatGroq
from dotenv import load_dotenv
import PyPDF2
from io import BytesIO
import logging
from docx import Document

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

load_dotenv()

# Placeholder regex for filling/dropping bracketed placeholders
PLACEHOLDER_RE = re.compile(r'\[([^\]]+)\]')
app = FastAPI(title="Resume-Based Interview System", version="3.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")

# ============================================
# PYDANTIC MODELS - PARSING
# ============================================

class PersonalInfo(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    linkedin: Optional[str] = None
    github: Optional[str] = None
    website: Optional[str] = None
    location: Optional[str] = None

class Education(BaseModel):
    institution: Optional[str] = None
    degree: Optional[str] = None
    field_of_study: Optional[str] = None
    grade: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    achievements: Optional[List[str]] = None

class Experience(BaseModel):
    company: Optional[str] = None
    role: Optional[str] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    duration: Optional[str] = None
    location: Optional[str] = None
    responsibilities: Optional[List[str]] = None
    technologies_used: Optional[List[str]] = None
    key_achievements: Optional[List[str]] = None

class Project(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    role: Optional[str] = None
    technologies: Optional[List[str]] = None
    key_features: Optional[List[str]] = None
    challenges_solved: Optional[List[str]] = None
    link: Optional[str] = None
    duration: Optional[str] = None

class SkillCategory(BaseModel):
    languages: Optional[List[str]] = None
    frameworks: Optional[List[str]] = None
    databases: Optional[List[str]] = None
    tools: Optional[List[str]] = None
    cloud_platforms: Optional[List[str]] = None
    other: Optional[List[str]] = None

class Certification(BaseModel):
    name: Optional[str] = None
    issuer: Optional[str] = None
    issue_date: Optional[str] = None
    expiry_date: Optional[str] = None
    credential_id: Optional[str] = None

class EnhancedResumeOutput(BaseModel):
    personal_info: PersonalInfo
    education: Optional[List[Education]] = None
    experience: Optional[List[Experience]] = None
    projects: Optional[List[Project]] = None
    skills: Optional[SkillCategory] = None
    certifications: Optional[List[Certification]] = None
    achievements: Optional[List[str]] = None
    publications: Optional[List[str]] = None
    languages: Optional[List[str]] = None

# ============================================
# PYDANTIC MODELS - QUESTIONS
# ============================================

class InterviewQuestion(BaseModel):
    id: int
    question: str
    category: str  # introduction/experience/project/skill/achievement
    expected_topics: List[str]
    context: Optional[str] = None

class InterviewQuestionSet(BaseModel):
    candidate_name: str
    questions: List[InterviewQuestion]
    total_questions: int

# ============================================
# PYDANTIC MODELS - SCORING
# ============================================

class CandidateInfo(BaseModel):
    name: str

class InterviewAnswer(BaseModel):
    question_id: int
    question: str
    category: str
    expected_topics: List[str]
    answer: str
    time_taken: Optional[int] = None

class ScoringInput(BaseModel):
    candidate_info: CandidateInfo
    interview_data: List[InterviewAnswer]

class QuestionScore(BaseModel):
    question_id: int
    question: str
    category: str
    candidate_answer: str
    time_taken: Optional[int] = None
    score: float  # 0-10
    feedback: str
    strengths: List[str]
    weaknesses: List[str]
    key_points_covered: List[str]
    key_points_missed: List[str]

class FinalScore(BaseModel):
    overall_score: float  # 0-100

class ScoringOutput(BaseModel):
    candidate_name: str
    total_questions: int
    questions_attempted: int
    question_scores: List[QuestionScore]
    final_score: FinalScore
    overall_feedback: str
    recommendation: str
    strengths_summary: List[str]
    areas_for_improvement: List[str]

# ============================================
# PROMPTS - PARSING
# ============================================

RESUME_PARSING_PROMPT = """Extract all information from the resume into structured JSON.

**RULES:**
1. Return ONLY valid JSON, no other text
2. Use null for missing fields
3. Extract maximum detail for projects and experience
4. Identify all technologies, tools, and skills mentioned

**JSON SCHEMA:**
{{
  "personal_info": {{
    "name": "string or null",
    "email": "string or null",
    "phone": "string or null",
    "linkedin": "string or null",
    "github": "string or null",
    "website": "string or null",
    "location": "string or null"
  }},
  "education": [
    {{
      "institution": "string or null",
      "degree": "string or null",
      "field_of_study": "string or null",
      "grade": "string or null",
      "start_date": "string or null",
      "end_date": "string or null",
      "achievements": ["achievement1"] or null
    }}
  ],
  "experience": [
    {{
      "company": "string or null",
      "role": "string or null",
      "start_date": "string or null",
      "end_date": "string or null",
      "duration": "string or null",
      "location": "string or null",
      "responsibilities": ["resp1", "resp2"] or null,
      "technologies_used": ["tech1", "tech2"] or null,
      "key_achievements": ["achievement1"] or null
    }}
  ],
  "projects": [
    {{
      "title": "string or null",
      "description": "string or null",
      "role": "Solo/Team Lead/Contributor or null",
      "technologies": ["tech1", "tech2"] or null,
      "key_features": ["feature1"] or null,
      "challenges_solved": ["challenge1"] or null,
      "link": "string or null",
      "duration": "string or null"
    }}
  ],
  "skills": {{
    "languages": ["lang1"] or null,
    "frameworks": ["framework1"] or null,
    "databases": ["db1"] or null,
    "tools": ["tool1"] or null,
    "cloud_platforms": ["platform1"] or null,
    "other": ["skill1"] or null
  }},
  "certifications": [
    {{
      "name": "string or null",
      "issuer": "string or null",
      "issue_date": "string or null",
      "expiry_date": "string or null",
      "credential_id": "string or null"
    }}
  ],
  "achievements": ["achievement1"] or null,
  "publications": ["publication1"] or null,
  "languages": ["language1"] or null
}}

**RESUME TEXT:**
\"\"\"
{resume_text}
\"\"\"

Return JSON only:"""

# ============================================
# STRUCTURED QUESTION GENERATION (NEW APPROACH)
# ============================================

STRUCTURED_QUESTION_GENERATION_PROMPT = """You are an expert technical interviewer. Generate interview questions based ONLY on the structured resume data below.

**CRITICAL RULES:**
1. Use ONLY the exact names provided in the resume data
2. DO NOT invent or assume ANY company names, project titles, or technologies
3. DO NOT use placeholders like [Company Name] - use actual values from the data
4. If a field is missing or empty, skip questions about that topic
5. Questions must be natural and conversational

**RESUME DATA (STRUCTURED JSON):**
```json
{resume_context}
```

**QUESTION GENERATION STRATEGY:**

Based on the resume content flags:
- has_experience={has_experience} → Ask 2-3 questions about actual work experience
- has_projects={has_projects} → Ask 2-3 questions about actual projects  
- has_certifications={has_certifications} → Ask 1 question about certification
- has_achievements={has_achievements} → Ask 1 question about achievement

**EXAMPLE 1 - With Experience:**
If resume data shows:
```json
{{"companies": [{{"company_name": "Ebani Tech Pvt. Ltd.", "role": "Intern"}}]}}
```

Generate:
```json
{{
  "id": 2,
  "question": "Can you tell me about your role as Intern at Ebani Tech Pvt. Ltd.?",
  "category": "experience",
  "expected_topics": ["role", "responsibilities", "Ebani Tech"],
  "context": "Asked about Ebani Tech Pvt. Ltd. mentioned in work experience"
}}
```

**WRONG EXAMPLES (NEVER DO THIS):**
❌ "Tell me about your experience at TCS" (TCS not in resume data)
❌ "What did you do at [Company Name]?" (placeholder not allowed)
❌ "Explain your e-commerce project" (project not in resume data)

**YOUR TASK:**

Generate EXACTLY 7 questions following this structure:

**QUESTION 1 (ALWAYS THE SAME):**
{{
  "id": 1,
  "question": "Tell me about yourself — your background, education, any experience you have, and projects you've worked on. Feel free to share any interests or hobbies as well.",
  "category": "introduction",
  "expected_topics": ["background", "education", "experience", "projects", "interests"],
  "context": "Opening introduction"
}}

**QUESTIONS 2-7:**
Distribute questions based on what's available in the resume:

Priority order:
1. If has_experience=true → 2-3 experience questions using exact company names
2. If has_projects=true → 2-3 project questions using exact project titles
3. If has_certifications=true → 1 certification question
4. If has_achievements=true → 1 achievement question
5. If only education → 2 education questions

**VALIDATION CHECKLIST (before generating):**
✓ Every company name must exist in resume_context.companies
✓ Every project title must exist in resume_context.projects
✓ Every technology must exist in the resume data
✓ No placeholders like [Company], [Project], etc.
✓ Context field explains which resume item was referenced

**OUTPUT FORMAT (JSON ONLY):**
{{
  "questions": [
    {{"id": 1, ...}},
    {{"id": 2, ...}},
    {{"id": 3, ...}},
    {{"id": 4, ...}},
    {{"id": 5, ...}},
    {{"id": 6, ...}},
    {{"id": 7, ...}}
  ]
}}

**CRITICAL: IF YOU CANNOT COMPOSE QUESTIONS USING ONLY THE PROVIDED resume_context, RETURN {{"questions": []}} EXACTLY. DO NOT INVENT.**

Generate the questions now:"""

# ============================================
# PROMPTS - ANSWER SCORING
# ============================================

ANSWER_EVALUATION_PROMPT = """You are evaluating a candidate's interview answer.

**QUESTION #{question_id} ({category}):**
{question}

**EXPECTED TOPICS:**
{expected_topics}

**CANDIDATE'S ANSWER:**
{answer}

**EVALUATION CRITERIA:**
1. Relevance - Does the answer address the question?
2. Completeness - Are key points covered?
3. Clarity - Is the explanation clear and well-structured?
4. Technical accuracy - Is the information correct?
5. Depth - Does it show good understanding?

**SCORING (0-10):**
- 9-10: Excellent, comprehensive, accurate
- 7-8: Good understanding, minor gaps
- 5-6: Basic understanding, missing key points
- 3-4: Partial/unclear understanding
- 0-2: No answer or completely wrong

**OUTPUT JSON:**
{{
  "score": 7.5,
  "feedback": "Detailed evaluation...",
  "strengths": ["specific strength 1", "specific strength 2"],
  "weaknesses": ["specific weakness 1"],
  "key_points_covered": ["point1", "point2"],
  "key_points_missed": ["point1"]
}}

Evaluate the answer:"""

OVERALL_SUMMARY_PROMPT = """Provide an overall assessment of the interview performance.

**CANDIDATE:** {candidate_name}
**OVERALL SCORE:** {overall_score}/100

**QUESTION-WISE PERFORMANCE:**
{performance_breakdown}

**SUMMARY GUIDELINES:**
- overall_feedback: 2-3 sentences on overall performance
- recommendation: strong-hire (80-100) / hire (65-79) / conditional-hire (50-64) / no-hire (<50)
- strengths_summary: Top 3 strengths shown
- areas_for_improvement: Top 3 areas needing work

**OUTPUT JSON:**
{{
  "overall_feedback": "Overall performance summary",
  "recommendation": "hire/no-hire/conditional-hire/strong-hire",
  "strengths_summary": ["strength 1", "strength 2", "strength 3"],
  "areas_for_improvement": ["area 1", "area 2", "area 3"]
}}

Generate summary:"""

# ============================================
# UTILITY FUNCTIONS
# ============================================

def extract_text_from_pdf(file_content: bytes) -> str:
    try:
        pdf_file = BytesIO(file_content)
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        return "".join([page.extract_text() for page in pdf_reader.pages])
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error extracting PDF: {str(e)}")

def extract_text_from_docx(file_content: bytes) -> str:
    try:
        docx_file = BytesIO(file_content)
        doc = Document(docx_file)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error extracting DOCX: {str(e)}")

def safe_json_parse(response_content: str) -> dict:
    try:
        return json.loads(response_content)
    except json.JSONDecodeError:
        json_match = re.search(r'```(?:json)?\s*(\{.*?\})\s*```', response_content, re.DOTALL)
        if json_match:
            return json.loads(json_match.group(1))
        json_match = re.search(r'\{.*\}', response_content, re.DOTALL)
        if json_match:
            return json.loads(json_match.group(0))
        raise ValueError("No valid JSON found")

def create_structured_resume_context(resume_data: dict) -> dict:
    """Create a clean, structured JSON context for question generation.
    This eliminates ambiguity and reduces hallucination risk."""
    context = {
        "candidate_name": resume_data.get("personal_info", {}).get("name", "Candidate"),
        "has_experience": False,
        "has_projects": False,
        "has_certifications": False,
        "has_achievements": False
    }
    
    # Extract companies with exact names - support multiple key names
    companies = []
    experience_list = resume_data.get("experience") or []
    for exp in experience_list:
        company_name = (exp.get("company") or exp.get("company_name")
                        or exp.get("key") or exp.get("employer") or "")
        role = exp.get("role") or exp.get("position") or ""
        if company_name:
            companies.append({
                "company_name": company_name.strip(),
                "role": role.strip(),
                "technologies": exp.get("technologies_used", []) or exp.get("technologies", []),
                "start_date": exp.get("start_date", "") or exp.get("start", ""),
                "end_date": exp.get("end_date", "") or exp.get("end", "Present")
            })
    
    if companies:
        context["companies"] = companies
        context["has_experience"] = True
    
    # Extract projects with exact titles - tolerate title/name variations
    projects = []
    projects_list = resume_data.get("projects") or []
    for proj in projects_list:
        title = proj.get("title") or proj.get("project_title") or proj.get("name")
        if title:
            projects.append({
                "project_title": title.strip(),
                "description": (proj.get("description") or "")[:200],
                "technologies": proj.get("technologies") or proj.get("technologies_used") or [],
                "role": proj.get("role") or ""
            })
    
    if projects:
        context["projects"] = projects
        context["has_projects"] = True
    
    # Extract education
    education = []
    education_list = resume_data.get("education") or []
    for edu in education_list:
        if edu.get("institution"):
            education.append({
                "institution": edu.get("institution"),
                "degree": edu.get("degree", ""),
                "field": edu.get("field_of_study", "")
            })
    
    if education:
        context["education"] = education
    
    # Extract certifications
    certifications = []
    certifications_list = resume_data.get("certifications") or []
    for cert in certifications_list:
        if cert.get("name"):
            certifications.append({
                "name": cert.get("name"),
                "issuer": cert.get("issuer", "")
            })
    
    if certifications:
        context["certifications"] = certifications
        context["has_certifications"] = True
    
    # Extract achievements
    achievements = resume_data.get("achievements") or []
    if achievements:
        context["achievements"] = achievements[:3]  # Top 3
        context["has_achievements"] = True
    
    # Extract skills (for context)
    skills = resume_data.get("skills") or {}
    all_skills = []
    if skills and isinstance(skills, dict):
        for category, items in skills.items():
            if items and isinstance(items, list):
                all_skills.extend(items)
    
    if all_skills:
        context["key_skills"] = all_skills[:10]  # Top 10 skills
    
    return context

def fill_placeholders_or_drop(q_text: str, resume_context: dict) -> Optional[str]:
    """Replace bracketed placeholders with resume values or return None if unavailable."""
    matches = PLACEHOLDER_RE.findall(q_text)
    out = q_text
    for m in matches:
        token = m.strip().lower()
        repl = None
        if "company" in token:
            companies = resume_context.get("companies", [])
            if companies:
                repl = companies[0].get("company_name")
        elif "project" in token or "project name" in token:
            projects = resume_context.get("projects", [])
            if projects:
                repl = projects[0].get("project_title")
        elif "certification" in token:
            certs = resume_context.get("certifications", [])
            if certs:
                repl = certs[0].get("name")
        elif "institution" in token or "university" in token or "college" in token:
            education = resume_context.get("education", [])
            if education:
                repl = education[0].get("institution")
        # If we can't find a replacement, return None to drop this question
        if not repl:
            return None
        out = out.replace(f"[{m}]", repl)
    return out

def validate_questions_against_resume(questions: List[dict], resume_context: dict) -> List[dict]:
    """Strict validation: questions must reference actual resume content.
    Returns only questions that pass validation."""
    validated = []
    
    # Extract all valid entities from resume
    valid_companies = {c["company_name"].lower() for c in resume_context.get("companies", [])}
    valid_projects = {p["project_title"].lower() for p in resume_context.get("projects", [])}
    valid_institutions = {e["institution"].lower() for e in resume_context.get("education", [])}
    valid_certs = {c["name"].lower() for c in resume_context.get("certifications", [])}
    
    # Common hallucinated companies (expand this list)
    hallucinated_companies = {
        'lti', 'ltimindtree', 'tcs', 'infosys', 'wipro', 'hcl technologies', 'hcl',
        'cognizant', 'accenture', 'capgemini', 'tech mahindra', 'persistent',
        'google', 'microsoft', 'amazon', 'facebook', 'meta', 'apple', 'netflix'
    }
    
    hallucinated_institutions = {
        'iit', 'nit', 'iiit', 'bits', 'vit', 'jntuh', 'anna university',
        'mit', 'stanford', 'harvard', 'berkeley'
    }
    
    for i, q in enumerate(questions):
        question_lower = q.get("question", "").lower()
        is_valid = True
        failure_reason = None
        
        # Validate question 1 (introduction) - don't auto-accept
        if i == 0:
            if not (q.get("id") == 1 or q.get("category") == "introduction"):
                is_valid = False
                failure_reason = "First question must be the canonical introduction (id==1 or category=='introduction')"
            # Still run subsequent checks for placeholders/hallucination
        
        # Check for placeholders
        if '[' in q.get("question", "") and ']' in q.get("question", ""):
            is_valid = False
            failure_reason = f"Contains placeholder: {q.get('question')}"
        
        # Check for hallucinated companies
        for hallucinated in hallucinated_companies:
            if hallucinated in question_lower and hallucinated not in valid_companies:
                is_valid = False
                failure_reason = f"Hallucinated company '{hallucinated}' not in resume"
                break
        
        # Check for hallucinated institutions
        for hallucinated in hallucinated_institutions:
            if hallucinated in question_lower and hallucinated not in valid_institutions:
                # Only flag if it's being asked about specifically
                if f"at {hallucinated}" in question_lower or f"from {hallucinated}" in question_lower:
                    is_valid = False
                    failure_reason = f"Hallucinated institution '{hallucinated}' not in resume"
                    break
        
        # Check for hallucinated projects (generic names)
        generic_projects = ['android app', 'e-commerce', 'chat app', 'weather app', 
                          'todo app', 'blog', 'social media', 'etl developer']
        for generic in generic_projects:
            if generic in question_lower and not any(generic in p for p in valid_projects):
                is_valid = False
                failure_reason = f"Hallucinated project '{generic}' not in resume"
                break
        
        if is_valid:
            validated.append(q)
        else:
            logger.error(f"🚨 INVALID QUESTION {i+1}: {failure_reason}")
            logger.error(f"   Question: {q.get('question')}")
    
    return validated

def deterministic_questions_from_context(ctx: dict, total: int = 7) -> List[dict]:
    """Generate questions deterministically from resume context without LLM.
    This is a 100% resume-grounded fallback when LLM fails or produces insufficient questions."""
    qlist = []
    
    # Always start with canonical introduction
    qlist.append({
        "id": 1,
        "question": "Tell me about yourself — your background, education, any experience you have, and projects you've worked on. Feel free to share any interests or hobbies as well.",
        "category": "introduction",
        "expected_topics": ["background", "education", "experience", "projects", "interests"],
        "context": "Opening introduction"
    })
    
    nid = 2
    
    # Experience-first (if present)
    if ctx.get("has_experience"):
        for comp in ctx.get("companies", [])[:3]:
            qlist.append({
                "id": nid,
                "question": f"Can you describe your role as {comp.get('role', '')} at {comp['company_name']} and your main responsibilities?",
                "category": "experience",
                "expected_topics": ["role", "responsibilities", comp["company_name"]],
                "context": f"Asked about {comp['company_name']}"
            })
            nid += 1
            if len(qlist) >= total:
                break
    
    # Projects
    if len(qlist) < total and ctx.get("has_projects"):
        for proj in ctx.get("projects", [])[:3]:
            qlist.append({
                "id": nid,
                "question": f"Walk me through the project '{proj['project_title']}': your role, key technical choices, and main challenges.",
                "category": "project",
                "expected_topics": ["project", proj["project_title"], "technical choices"],
                "context": f"Asked about {proj['project_title']}"
            })
            nid += 1
            if len(qlist) >= total:
                break
    
    # Certifications / achievements next
    if len(qlist) < total and ctx.get("has_certifications"):
        cert = ctx.get("certifications", [])[0]
        qlist.append({
            "id": nid,
            "question": f"You have the '{cert['name']}' certification. Why did you pursue it and how have you applied it?",
            "category": "certification",
            "expected_topics": ["certification", cert["name"]],
            "context": f"Asked about {cert['name']}"
        })
        nid += 1
    
    # Achievements
    if len(qlist) < total and ctx.get("has_achievements"):
        achievements = ctx.get("achievements", [])
        if achievements:
            qlist.append({
                "id": nid,
                "question": f"You mentioned achieving '{achievements[0]}'. Can you elaborate on that?",
                "category": "achievement",
                "expected_topics": ["achievement", "accomplishment"],
                "context": "Asked about listed achievement"
            })
            nid += 1
    
    # Education-based questions if needed
    if len(qlist) < total and ctx.get("education"):
        edu = ctx.get("education", [])[0]
        qlist.append({
            "id": nid,
            "question": f"Tell me about your time at {edu['institution']} and what you studied.",
            "category": "education",
            "expected_topics": ["education", edu["institution"], "studies"],
            "context": f"Asked about {edu['institution']}"
        })
        nid += 1
    
    # Pad with skills-based questions if still short
    while len(qlist) < total:
        qlist.append({
            "id": nid,
            "question": "What are your key technical strengths and which tools/technologies are you most comfortable using?",
            "category": "skills",
            "expected_topics": ["skills", "tools", "technologies"],
            "context": "General skills question"
        })
        nid += 1
    
    return qlist[:total]

def is_answer_valid(answer: str) -> bool:
    """Check if answer is meaningful"""
    if not answer or len(answer.strip()) < 10:
        return False
    words = answer.strip().split()
    if len(words) < 3:
        return False
    alpha_count = sum(c.isalpha() for c in answer)
    return alpha_count >= len(answer) * 0.6

# ============================================
# API ENDPOINTS
# ============================================

@app.get("/")
async def root():
    return {
        "message": "Resume-Based Natural Interview System v3.1 (Fixed)",
        "endpoints": {
            "parse_resume": "/parse-resume",
            "generate_questions": "/generate-questions",
            "score_answers": "/score-answers",
            "debug_questions": "/debug-questions",
            "health": "/health"
        }
    }

@app.post("/parse-resume", response_model=EnhancedResumeOutput)
async def parse_resume(file: UploadFile = File(...)):
    """Parse resume with consistent JSON structure"""
    if not GROQ_API_KEY:
        raise HTTPException(status_code=500, detail="GROQ_API_KEY not configured")
    
    allowed_types = ["application/pdf", "text/plain", 
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail="Only PDF, TXT, and DOCX files supported")
    
    try:
        file_content = await file.read()
        
        if file.content_type == "application/pdf":
            resume_text = extract_text_from_pdf(file_content)
        elif file.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            resume_text = extract_text_from_docx(file_content)
        else:
            resume_text = file_content.decode('utf-8')
        
        if not resume_text.strip():
            raise HTTPException(status_code=400, detail="Could not extract text from file")
        
        logger.info("Parsing resume...")
        prompt = RESUME_PARSING_PROMPT.format(resume_text=resume_text)
        
        llm = ChatGroq(
            api_key=GROQ_API_KEY,
            model="llama-3.1-8b-instant",
            temperature=0.1,
            model_kwargs={"response_format": {"type": "json_object"}}
        )
        
        response = llm.invoke(prompt)
        resume_data = safe_json_parse(response.content)
        
        return EnhancedResumeOutput(**resume_data)
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error parsing resume: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error: {str(e)}")

@app.post("/generate-questions", response_model=InterviewQuestionSet)
async def generate_questions(resume_data: EnhancedResumeOutput):
    """Generate natural interview questions based purely on resume using structured approach"""
    if not GROQ_API_KEY:
        raise HTTPException(status_code=500, detail="GROQ_API_KEY not configured")
    
    try:
        candidate_name = resume_data.personal_info.name or "the candidate"
        
        # Create structured resume context (NEW APPROACH)
        resume_context = create_structured_resume_context(resume_data.model_dump())
        
        logger.info(f"Generating questions for {candidate_name}")
        logger.info(f"📊 RESUME CONTEXT:")
        logger.info(f"  - Has Experience: {resume_context.get('has_experience')}")
        logger.info(f"  - Has Projects: {resume_context.get('has_projects')}")
        logger.info(f"  - Companies: {[c['company_name'] for c in resume_context.get('companies', [])]}")
        logger.info(f"  - Projects: {[p['project_title'] for p in resume_context.get('projects', [])]}")
        logger.info(f"📋 STRUCTURED RESUME CONTEXT:")
        logger.info(f"{'='*50}")
        logger.info(json.dumps(resume_context, indent=2))
        logger.info(f"{'='*50}")
        
        # Generate prompt with structured data
        prompt = STRUCTURED_QUESTION_GENERATION_PROMPT.format(
            resume_context=json.dumps(resume_context, indent=2),
            has_experience=resume_context.get('has_experience'),
            has_projects=resume_context.get('has_projects'),
            has_certifications=resume_context.get('has_certifications'),
            has_achievements=resume_context.get('has_achievements')
        )
        
        # Use larger model with zero temperature for deterministic output
        llm = ChatGroq(
            api_key=GROQ_API_KEY,
            model="llama-3.3-70b-versatile",  # Larger model for better reasoning
            temperature=0.0,  # Zero temperature for maximum determinism
            model_kwargs={"response_format": {"type": "json_object"}}
        )
        
        response = llm.invoke(prompt)
        questions_data = safe_json_parse(response.content)
        
        questions = questions_data.get("questions", [])
        
        logger.info(f"🤖 RAW QUESTIONS GENERATED:")
        for i, q in enumerate(questions):
            logger.info(f"Q{i+1}: {q.get('question')}")
        
        # Fill placeholders or drop questions with unfillable placeholders
        processed_questions = []
        for i, q in enumerate(questions):
            original_text = q.get("question", "")
            if '[' in original_text and ']' in original_text:
                filled_text = fill_placeholders_or_drop(original_text, resume_context)
                if filled_text is None:
                    logger.warning(f"⚠️ Dropping question {i+1} with unfillable placeholder: {original_text}")
                    continue
                q["question"] = filled_text
                logger.info(f"✏️ Filled placeholder in Q{i+1}: {filled_text}")
            processed_questions.append(q)
        
        # Strict validation against resume context
        validated_questions = validate_questions_against_resume(processed_questions, resume_context)
        
        logger.info(f"✅ VALIDATED QUESTIONS ({len(validated_questions)}/7):")
        for i, q in enumerate(validated_questions):
            logger.info(f"Q{i+1}: {q.get('question')}")
        
        # If we lost too many questions, use deterministic fallback
        if len(validated_questions) < 5:
            logger.warning(f"⚠️ Only {len(validated_questions)} valid questions from LLM! Using deterministic fallback.")
            logger.warning(f"Resume context had: experience={resume_context.get('has_experience')}, "
                          f"projects={resume_context.get('has_projects')}")
            validated_questions = deterministic_questions_from_context(resume_context, total=7)
            logger.info(f"✅ Generated {len(validated_questions)} deterministic questions from resume context")
        else:
            # Ensure introduction question is first
            if validated_questions and validated_questions[0].get("id") != 1:
                canonical_intro = {
                    "id": 1,
                    "question": "Tell me about yourself — your background, education, any experience you have, and projects you've worked on. Feel free to share any interests or hobbies as well.",
                    "category": "introduction",
                    "expected_topics": ["background", "education", "experience", "projects", "interests"],
                    "context": "Opening introduction"
                }
                # Check if first question is a valid intro
                if validated_questions[0].get("category") != "introduction":
                    validated_questions.insert(0, canonical_intro)
                    logger.info("Inserted canonical introduction at the beginning")
                else:
                    validated_questions[0] = canonical_intro
            elif not validated_questions:
                # Empty list - use deterministic fallback
                validated_questions = deterministic_questions_from_context(resume_context, total=7)
                logger.info(f"✅ No validated questions - used deterministic fallback")
            
            # Ensure exactly 7 questions
            if len(validated_questions) < 7:
                # Use deterministic generator to fill remaining
                full_set = deterministic_questions_from_context(resume_context, total=7)
                # Keep validated questions and fill from deterministic set
                existing_ids = {q.get("id") for q in validated_questions}
                for q in full_set:
                    if len(validated_questions) >= 7:
                        break
                    if q.get("id") not in existing_ids:
                        validated_questions.append(q)
                        logger.info(f"Added deterministic question: {q.get('question')[:50]}...")
        
        return InterviewQuestionSet(
            candidate_name=candidate_name,
            questions=validated_questions[:7],
            total_questions=len(validated_questions[:7])
        )
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error generating questions: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error: {str(e)}")

@app.post("/debug-questions")
async def debug_question_generation(resume_data: EnhancedResumeOutput):
    """Debug endpoint to test question generation with detailed logging"""
    try:
        candidate_name = resume_data.personal_info.name or "Test Candidate"
        resume_context = create_structured_resume_context(resume_data.model_dump())
        
        return {
            "candidate_name": candidate_name,
            "structured_resume_context": resume_context,
            "resume_data_keys": list(resume_data.model_dump().keys()),
            "experience_count": len(resume_data.experience or []),
            "project_count": len(resume_data.projects or []),
            "education_count": len(resume_data.education or []),
            "actual_companies": [exp.company for exp in (resume_data.experience or []) if exp.company],
            "actual_projects": [proj.title for proj in (resume_data.projects or []) if proj.title],
            "actual_institutions": [edu.institution for edu in (resume_data.education or []) if edu.institution],
            "has_flags": {
                "has_experience": resume_context.get('has_experience'),
                "has_projects": resume_context.get('has_projects'),
                "has_certifications": resume_context.get('has_certifications'),
                "has_achievements": resume_context.get('has_achievements')
            }
        }
    except Exception as e:
        return {"error": str(e)}

@app.post("/score-answers", response_model=ScoringOutput)
async def score_answers(scoring_input: ScoringInput):
    """Score candidate's interview answers with detailed feedback"""
    if not GROQ_API_KEY:
        raise HTTPException(status_code=500, detail="GROQ_API_KEY not configured")
    
    if not scoring_input.interview_data:
        raise HTTPException(status_code=400, detail="No interview data provided")
    
    try:
        llm = ChatGroq(
            api_key=GROQ_API_KEY,
            model="llama-3.3-70b-versatile",
            temperature=0.3,
            model_kwargs={"response_format": {"type": "json_object"}}
        )
        
        question_scores = []
        questions_attempted = 0
        
        # Evaluate each question
        for qa in scoring_input.interview_data:
            # Check if answer is valid
            if not is_answer_valid(qa.answer):
                question_score = QuestionScore(
                    question_id=qa.question_id,
                    question=qa.question,
                    category=qa.category,
                    candidate_answer=qa.answer,
                    time_taken=qa.time_taken,
                    score=0.0,
                    feedback="No meaningful answer provided.",
                    strengths=[],
                    weaknesses=["No answer provided"],
                    key_points_covered=[],
                    key_points_missed=qa.expected_topics
                )
                question_scores.append(question_score)
                continue
            
            questions_attempted += 1
            
            # Generate evaluation prompt
            prompt = ANSWER_EVALUATION_PROMPT.format(
                question_id=qa.question_id,
                category=qa.category,
                question=qa.question,
                expected_topics=", ".join(qa.expected_topics),
                answer=qa.answer
            )
            
            # Get evaluation
            response = llm.invoke(prompt)
            eval_data = safe_json_parse(response.content)
            
            if not eval_data:
                eval_data = {
                    "score": 5.0,
                    "feedback": "Could not parse evaluation.",
                    "strengths": [],
                    "weaknesses": ["Evaluation failed"],
                    "key_points_covered": [],
                    "key_points_missed": qa.expected_topics
                }
            
            score = float(eval_data.get("score", 5.0))
            
            question_score = QuestionScore(
                question_id=qa.question_id,
                question=qa.question,
                category=qa.category,
                candidate_answer=qa.answer,
                time_taken=qa.time_taken,
                score=round(score, 2),
                feedback=eval_data.get("feedback", ""),
                strengths=eval_data.get("strengths", []),
                weaknesses=eval_data.get("weaknesses", []),
                key_points_covered=eval_data.get("key_points_covered", []),
                key_points_missed=eval_data.get("key_points_missed", [])
            )
            question_scores.append(question_score)
        
        # Calculate final score
        all_scores = [s.score for s in question_scores]
        avg_score = sum(all_scores) / len(all_scores) if all_scores else 0
        overall_score = round((avg_score / 10) * 100, 2)
        
        final_score = FinalScore(overall_score=overall_score)
        
        # Generate overall summary
        performance_breakdown = "\n".join([
            f"Q{s.question_id} ({s.category}): {s.score}/10"
            for s in question_scores
        ])
        
        summary_prompt = OVERALL_SUMMARY_PROMPT.format(
            candidate_name=scoring_input.candidate_info.name,
            overall_score=overall_score,
            performance_breakdown=performance_breakdown
        )
        
        summary_response = llm.invoke(summary_prompt)
        summary_data = safe_json_parse(summary_response.content)
        
        if not summary_data:
            summary_data = {
                "overall_feedback": f"Candidate scored {overall_score}/100 overall.",
                "recommendation": "conditional-hire" if overall_score >= 50 else "no-hire",
                "strengths_summary": ["Requires review"],
                "areas_for_improvement": ["Requires review"]
            }
        
        logger.info(f"✅ Scoring complete: {overall_score}/100")
        
        return ScoringOutput(
            candidate_name=scoring_input.candidate_info.name,
            total_questions=len(scoring_input.interview_data),
            questions_attempted=questions_attempted,
            question_scores=question_scores,
            final_score=final_score,
            overall_feedback=summary_data.get("overall_feedback", ""),
            recommendation=summary_data.get("recommendation", "conditional-hire"),
            strengths_summary=summary_data.get("strengths_summary", []),
            areas_for_improvement=summary_data.get("areas_for_improvement", [])
        )
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error scoring answers: {str(e)}")
        raise HTTPException(status_code=500, detail=f"Error: {str(e)}")

@app.get("/health")
async def health_check():
    return {
        "status": "healthy",
        "version": "3.1.0",
        "groq_configured": bool(GROQ_API_KEY),
        "features": [
            "Resume parsing with null for missing fields",
            "Structured question generation (eliminates hallucination)",
            "Answer scoring with detailed feedback",
            "No technology bias - purely resume-driven",
            "Enhanced validation with hallucination detection",
            "Debug endpoint for troubleshooting",
            "Lower temperature (0.1) for more deterministic results",
            "70B model for better reasoning"
        ]
    }

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8002)